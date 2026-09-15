"""Weekly field-rep pipeline chase lists -> the owner's OneDrive `rep_pipelines` folder.

Rebuilds, IN PLACE, the workbooks first produced by hand on 20 Aug 2026:
  * PIPELINE_MASTER.xlsx            Summary sheet + one sheet per rep
  * pipeline_<Rep_Name>.xlsx        the rep's own chase list (shared to the rep)

Definition = the live metre's pipeline tile (lib/provider/bronze.ts): an
appointment that was attended in the last 60 days, chase status still
winnable, lead not dead, not sold, test/deleted leads out. Rep = latest active
booking's rep, else the CRM owner. "Sat this month" = attended appointments
dated this calendar month. Overdue = more than 14 days since the visit.

Files are overwritten through Microsoft Graph, so links and sharing survive
and OneDrive keeps version history. Anything a rep typed into the
Contacted? / Outcome / Notes columns is read back first and carried forward
for customers still on the list (owner, 15 Sep 2026: "edit the same one").

Usage:
  python -m scripts.rep_pipelines_weekly            # build + upload
  python -m scripts.rep_pipelines_weekly --dry-run  # build to ./out only
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from dotenv import load_dotenv
from google.cloud import bigquery
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ingestion.onedrive.auth import get_token  # noqa: E402
from shared.phone import normalise_phone  # noqa: E402

load_dotenv()
UK = ZoneInfo("Europe/London")
PROJECT = "trustwarehouse"
GRAPH = "https://graph.microsoft.com/v1.0"
FOLDER = os.environ.get("REP_PIPELINES_FOLDER", "rep_pipelines")
MASTER_NAME = "PIPELINE_MASTER.xlsx"
LEGACY_MASTER = "PIPELINE_MASTER_20_Aug.xlsx"  # renamed in place on first run
WINDOW_DAYS = 60
OVERDUE_AFTER_DAYS = 14
DEAD = ["Appointment Cancelled", "Not Interested", "Too Expensive", "Not a Lead", "Bought Elsewhere"]
CHASE = ["appointment sat", "follow up", "follow up text", "no contact", "not ready yet"]
HEADERS = ["Customer", "Phone 1", "Phone 2", "Phone 3", "Email", "Region", "Appointment date",
           "Days since visit", "Days overdue", "Status", "Contacted?", "Outcome", "Notes",
           "Quote £ (Unleashed)", "Lead ID"]
WIDTHS = [22, 15, 15, 15, 26, 14, 17, 14, 12, 16, 14, 20, 40, 16, 20]
CONTACTED = '"☐ Not yet,✅ Contacted"'
OUTCOMES = ('"No answer,Left message,Call back arranged,Appointment rebooked,SOLD,'
            'Not interested,Too expensive,Bought elsewhere,Other"')
HDR_FILL = PatternFill("solid", fgColor="1F2A44")
HDR_FONT = Font(bold=True, color="FFFFFF")


def log(msg: str) -> None:
    print(f"{datetime.now(UK):%H:%M:%S} {msg}", flush=True)


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------

def _bq() -> bigquery.Client:
    return bigquery.Client(project=PROJECT)


def sql_list(vals: list[str]) -> str:
    return ", ".join("'" + v.replace("'", "\\'") + "'" for v in vals)


def query_leads(bq: bigquery.Client) -> list[dict]:
    """Pipeline leads (metre definition) with contact details; plus this month's sits."""
    sql = f"""
    WITH latest AS (
      SELECT * FROM `{PROJECT}.bronze.sharpspring_leads`
      WHERE first_name NOT LIKE 'Zzz%'
        AND TRIM(CONCAT(first_name, ' ', COALESCE(last_name, ''))) NOT LIKE 'Test %'
        AND LOWER(COALESCE(first_name, '')) != 'test'
        AND CAST(id AS STRING) NOT IN (SELECT id FROM `{PROJECT}.bronze.sharpspring_leads_deleted`)
      QUALIFY ROW_NUMBER() OVER (PARTITION BY id ORDER BY update_timestamp DESC) = 1
    )
    SELECT CAST(id AS STRING) AS lead_id,
           TRIM(CONCAT(COALESCE(first_name,''), ' ', COALESCE(last_name,''))) AS customer,
           COALESCE(mobile_phone_number, '') AS phone1,
           COALESCE(phone_number, '') AS phone2,
           COALESCE(alternative_phone_number_5af46947e2fc1, '') AS phone3,
           COALESCE(email_address, '') AS email,
           COALESCE(location_6349396e4a08d, '') AS region,
           COALESCE(zipcode, '') AS postcode,
           FORMAT_TIMESTAMP('%Y-%m-%d %H:%M', appointment_time___date_5ae8ca2f532bc, 'Europe/London') AS appt,
           CAST(DATE(appointment_time___date_5ae8ca2f532bc, 'Europe/London') AS STRING) AS appt_date,
           CAST(owner_id AS STRING) AS owner_id,
           LOWER(TRIM(appointment_status_637f8d6fa1096)) AS chase_status,
           COALESCE(status_633ae6f6ac6fe, '') AS lead_status,
           (LOWER(TRIM(appointment_status_637f8d6fa1096)) IN ({sql_list(CHASE)})
             AND status_633ae6f6ac6fe NOT IN ({sql_list(DEAD)})
             AND DATE(appointment_time___date_5ae8ca2f532bc, 'Europe/London')
                 >= DATE_SUB(CURRENT_DATE('Europe/London'), INTERVAL {WINDOW_DAYS} DAY)) AS in_pipeline,
           (DATE_TRUNC(DATE(appointment_time___date_5ae8ca2f532bc, 'Europe/London'), MONTH)
             = DATE_TRUNC(CURRENT_DATE('Europe/London'), MONTH)
             AND TRIM(COALESCE(appointment_status_637f8d6fa1096, '')) != '') AS sat_this_month
    FROM latest
    WHERE appointment_time___date_5ae8ca2f532bc IS NOT NULL
      AND appointment_time___date_5ae8ca2f532bc <= CURRENT_TIMESTAMP()
      AND DATE(appointment_time___date_5ae8ca2f532bc, 'Europe/London')
          >= DATE_SUB(DATE_TRUNC(CURRENT_DATE('Europe/London'), MONTH), INTERVAL {WINDOW_DAYS} DAY)
    """
    rows = [dict(r) for r in bq.query(sql, location="europe-west2").result()]
    return [r for r in rows if r["in_pipeline"] or r["sat_this_month"]]


def query_reps(bq: bigquery.Client, lead_ids: list[str]) -> tuple[dict, dict, float]:
    """(booking_rep by lead, rep by CRM owner id, 2026 average sale £)."""
    sql = f"""
    SELECT 'booking' AS kind, lead_id AS k,
           ARRAY_AGG(rep_name ORDER BY booked_at DESC LIMIT 1)[OFFSET(0)] AS v
    FROM `{PROJECT}.app.bookings`
    WHERE status = 'active' AND rep_name IS NOT NULL AND rep_name != ''
      AND lead_id IN UNNEST(@ids) GROUP BY lead_id
    UNION ALL SELECT 'rep', name, CAST(sharpspring_owner_id AS STRING) FROM `{PROJECT}.app.reps`
    UNION ALL SELECT 'avg', CAST(ROUND(AVG(COALESCE(heating_amount,0)+COALESCE(water_amount,0)+COALESCE(chc_amount,0)),2) AS STRING), ''
    FROM `{PROJECT}.app.sales` WHERE status='active' AND customer_name NOT LIKE 'Zzz Testlead%'
      AND EXTRACT(YEAR FROM sale_date) = EXTRACT(YEAR FROM CURRENT_DATE('Europe/London'))
    """
    job = bq.query(sql, location="US", job_config=bigquery.QueryJobConfig(
        query_parameters=[bigquery.ArrayQueryParameter("ids", "STRING", lead_ids or [""])]))
    booking_rep, rep_by_owner, avg = {}, {}, 0.0
    for r in job.result():
        if r.kind == "booking" and r.v:
            booking_rep[r.k] = r.v
        elif r.kind == "rep" and r.v:
            rep_by_owner[r.v] = r.k
        elif r.kind == "avg":
            avg = float(r.k or 0)
    return booking_rep, rep_by_owner, avg


# ---------------------------------------------------------------------------
# Unleashed quotes (best effort — the list still ships if the API is down)
# ---------------------------------------------------------------------------

def _norm_name(s: str) -> str:
    toks = re.sub(r"[^a-z ]", "", (s or "").lower()).split()
    out: list[str] = []
    for t in toks:
        if not out or out[-1] != t:
            out.append(t)
    return " ".join(out)


def _norm_pc(s: str) -> str:
    return re.sub(r"\s+", "", (s or "").upper())


def unleashed_quotes(lookback_days: int = 130) -> list[dict]:
    from ingestion.unleashed.client import UnleashedClient
    c = UnleashedClient()
    since = (date.today() - timedelta(days=lookback_days)).strftime("%Y-%m-%dT00:00:00")
    rank = {"Accepted": 0, "Pending": 1, "Draft": 2, "Cancelled": 3}
    out = []
    for q in c.paginate("SalesQuotes", f"modifiedSince={since}"):
        cust = q.get("Customer") or {}
        m = re.search(r"/Date\((\d+)", q.get("QuoteDate") or "")
        out.append({
            "name": _norm_name(cust.get("CustomerName") or q.get("DeliveryName") or ""),
            "pc": _norm_pc(q.get("DeliveryPostCode")),
            "status": q.get("QuoteStatus"), "rank": rank.get(q.get("QuoteStatus"), 9),
            "date": datetime.fromtimestamp(int(m.group(1)) / 1000, timezone.utc).date().isoformat() if m else "",
            "subtotal": float(q.get("SubTotal") or 0),
        })
    return out


def best_quote(lead: dict, quotes: list[dict]) -> float | None:
    ln, lp = _norm_name(lead["customer"]), _norm_pc(lead["postcode"])
    surname = ln.split()[-1] if ln else ""
    hits = [q for q in quotes if (ln and q["name"] == ln)
            or (lp and q["pc"] == lp and surname and surname in q["name"])]
    if not hits:
        return None
    hits.sort(key=lambda q: (q["rank"], q["date"]))
    top = hits[0]["rank"]
    return max((q for q in hits if q["rank"] == top), key=lambda q: q["date"])["subtotal"]


# ---------------------------------------------------------------------------
# OneDrive
# ---------------------------------------------------------------------------

def graph_headers() -> dict:
    return {"Authorization": f"Bearer {get_token()}"}


def folder_items() -> dict[str, dict]:
    r = requests.get(f"{GRAPH}/me/drive/root:/{FOLDER}:/children", headers=graph_headers(),
                     params={"$select": "id,name,size", "$top": 500}, timeout=60)
    r.raise_for_status()
    return {i["name"]: i for i in r.json().get("value", []) if not i["name"].startswith("~$")}


def download(item_id: str) -> bytes:
    r = requests.get(f"{GRAPH}/me/drive/items/{item_id}/content", headers=graph_headers(), timeout=120)
    r.raise_for_status()
    return r.content


def upload(name: str, data: bytes, items: dict[str, dict]) -> str:
    """Overwrite in place when the file exists (same item id, same links); else create."""
    if name in items:
        url = f"{GRAPH}/me/drive/items/{items[name]['id']}/content"
    else:
        url = f"{GRAPH}/me/drive/root:/{FOLDER}/{name}:/content"
    r = requests.put(url, headers={**graph_headers(), "Content-Type":
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}, data=data, timeout=180)
    r.raise_for_status()
    return "replaced" if name in items else "created"


def rename(item_id: str, new_name: str) -> None:
    r = requests.patch(f"{GRAPH}/me/drive/items/{item_id}", headers={**graph_headers(), "Content-Type": "application/json"},
                       json={"name": new_name}, timeout=60)
    r.raise_for_status()


def read_rep_edits(data: bytes) -> dict[tuple[str, str], tuple]:
    """(customer, phone1) -> (Contacted?, Outcome, Notes) for rows the rep touched."""
    edits: dict[tuple[str, str], tuple] = {}
    try:
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    except Exception:  # noqa: BLE001 — a corrupt file must not stop the run
        return edits
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
        if row and row[0] == "Customer":
            header_row = i
            break
    if header_row is None:
        return edits
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or not row[0]:
            continue
        contacted, outcome, notes = (row[10] if len(row) > 10 else None), (row[11] if len(row) > 11 else None), (row[12] if len(row) > 12 else None)
        touched = bool(outcome) or bool(notes) or (contacted and "Not yet" not in str(contacted))
        if touched:
            edits[(str(row[0]).strip().lower(), str(row[1] or "").strip())] = (contacted, outcome, notes)
    return edits


# ---------------------------------------------------------------------------
# Workbooks
# ---------------------------------------------------------------------------

def _header(ws, row: int) -> None:
    for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w


def _rows(ws, first_row: int, leads: list[dict]) -> None:
    for i, l in enumerate(leads):
        r = first_row + i
        appt = datetime.strptime(l["appt"], "%Y-%m-%d %H:%M").strftime("%d/%m/%Y %H:%M")
        vals = [l["customer"], l["phone1"] or None, l["phone2"] or None, l["phone3"] or None, l["email"] or None,
                l["region"] or None, appt, l["age"], l["overdue"], l["chase_status"],
                l["contacted"] or "☐ Not yet", l["outcome"], l["notes"], l["quote"], l["lead_id"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v)
        if l["quote"] is not None:
            ws.cell(r, 14).number_format = "£#,##0"
    last = first_row + max(len(leads), 1) - 1
    dv1 = DataValidation(type="list", formula1=CONTACTED, allow_blank=True)
    dv2 = DataValidation(type="list", formula1=OUTCOMES, allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    dv1.add(f"K{first_row}:K{last}"); dv2.add(f"L{first_row}:L{last}")
    ws.auto_filter.ref = f"A{first_row - 1}:{get_column_letter(len(HEADERS))}{last}"


def build_rep_file(rep: str, leads: list[dict], stats: dict, generated: str) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = rep[:31]
    ws["A1"] = f"PIPELINE CHASE LIST — {rep}"; ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:M1")
    ws["A2"] = (f"Generated: {generated} (Europe/London)  |  Pipeline leads: {stats['n']}  |  "
                f"Sat this month: {stats['sat']}  |  Overdue: {stats['overdue']}  |  Est £: {stats['est']:.2f}  |  "
                f"Quoted £: {stats['quoted']:.2f}  |  No callable phone: {stats['nophone']}")
    ws.merge_cells("A2:M2")
    _header(ws, 4); _rows(ws, 5, leads); ws.freeze_panes = "A5"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_master(by_rep: dict[str, list[dict]], stats: dict[str, dict], generated: str) -> bytes:
    wb = Workbook(); s = wb.active; s.title = "Summary"
    cols = ["Rep", "Sat this month", "Pipeline leads", "Est £", "Quoted £ (Unleashed)", "Overdue count", "Oldest days waiting"]
    for c, (h, w) in enumerate(zip(cols, [22, 16, 15, 15, 20, 15, 20]), start=1):
        cell = s.cell(1, c, h); cell.font, cell.fill = HDR_FONT, HDR_FILL
        s.column_dimensions[get_column_letter(c)].width = w
    tot = defaultdict(float); oldest = 0
    for rep in sorted(by_rep, key=lambda n: (n == "Unattributed", n)):
        st = stats[rep]
        s.append([rep, st["sat"], st["n"], round(st["est"], 2), round(st["quoted"], 2), st["overdue"], st["oldest"]])
        for k in ("sat", "n", "est", "quoted", "overdue"):
            tot[k] += st[k]
        oldest = max(oldest, st["oldest"])
    s.append(["TOTAL", int(tot["sat"]), int(tot["n"]), round(tot["est"], 2), round(tot["quoted"], 2), int(tot["overdue"]), oldest])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    s.append([]); s.append([f"Generated {generated} (Europe/London). Rebuilt every Monday from the warehouse; "
                            f"same definition as the live metre's pipeline tile (attended in the last {WINDOW_DAYS} days, "
                            f"still winnable, not sold, not dead). Overdue = more than {OVERDUE_AFTER_DAYS} days since the visit. "
                            "Est £ = leads × this year's average sale; Quoted £ = the customer's Unleashed quote where one exists."])
    s.freeze_panes = "A2"
    for rep in sorted(by_rep, key=lambda n: (n == "Unattributed", n)):
        ws = wb.create_sheet(rep[:31]); _header(ws, 1); _rows(ws, 2, by_rep[rep]); ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def rep_filename(rep: str) -> str:
    return "pipeline_" + re.sub(r"[^A-Za-z0-9]+", "_", rep).strip("_") + ".xlsx"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="write workbooks to ./out instead of OneDrive")
    ap.add_argument("--no-quotes", action="store_true", help="skip the Unleashed quote lookup")
    args = ap.parse_args()

    now = datetime.now(UK); today = now.date()
    generated = now.strftime("%d/%m/%Y %H:%M")
    bq = _bq()
    leads = query_leads(bq)
    pipeline = [l for l in leads if l["in_pipeline"]]
    booking_rep, rep_by_owner, avg = query_reps(bq, [l["lead_id"] for l in leads])
    log(f"leads in window: {len(leads)} | pipeline: {len(pipeline)} | avg sale £{avg:,.0f}")

    quotes: list[dict] = []
    if not args.no_quotes:
        try:
            quotes = unleashed_quotes()
            log(f"unleashed quotes loaded: {len(quotes)}")
        except Exception as exc:  # noqa: BLE001
            log(f"unleashed quotes unavailable ({exc}) - continuing without")

    for l in leads:
        # Pack the three CRM phone fields into Phone 1..3, normalised to 44...
        # (the 20 Aug files' convention), so Phone 1 is always the first callable.
        phones: list[str] = []
        for raw in (l["phone1"], l["phone2"], l["phone3"]):
            p = normalise_phone(raw)
            if p and p not in phones:
                phones.append(p)
        l["phone1"], l["phone2"], l["phone3"] = (phones + ["", "", ""])[:3]
        l["rep"] = booking_rep.get(l["lead_id"]) or rep_by_owner.get(l["owner_id"] or "") or "Unattributed"
        d = date.fromisoformat(l["appt_date"]); l["age"] = (today - d).days
        l["overdue"] = max(0, l["age"] - OVERDUE_AFTER_DAYS)
        l["quote"] = best_quote(l, quotes) if quotes else None
        l["contacted"] = l["outcome"] = l["notes"] = None

    by_rep: dict[str, list[dict]] = defaultdict(list)
    for l in pipeline:
        by_rep[l["rep"]].append(l)
    for rep in by_rep:
        by_rep[rep].sort(key=lambda l: -l["age"])
    sat = defaultdict(int)
    for l in leads:
        if l["sat_this_month"]:
            sat[l["rep"]] += 1

    items = {} if args.dry_run else folder_items()
    if not args.dry_run and LEGACY_MASTER in items and MASTER_NAME not in items:
        rename(items[LEGACY_MASTER]["id"], MASTER_NAME)
        items = folder_items()
        log(f"renamed {LEGACY_MASTER} -> {MASTER_NAME} (same file, same link)")

    # carry the reps' own entries forward
    carried = 0
    for rep, mine in by_rep.items():
        name = rep_filename(rep)
        if name in items:
            edits = read_rep_edits(download(items[name]["id"]))
            for l in mine:
                e = edits.get((l["customer"].strip().lower(), l["phone1"].strip()))
                if e:
                    l["contacted"], l["outcome"], l["notes"] = e; carried += 1
    log(f"rep entries carried forward: {carried}")

    stats = {}
    for rep, mine in by_rep.items():
        stats[rep] = {"n": len(mine), "sat": sat.get(rep, 0), "overdue": sum(1 for l in mine if l["overdue"] > 0),
                      "oldest": max((l["age"] for l in mine), default=0), "est": len(mine) * avg,
                      "quoted": sum(l["quote"] or 0 for l in mine),
                      "nophone": sum(1 for l in mine if not (l["phone1"] or l["phone2"] or l["phone3"]))}

    out = Path(__file__).resolve().parents[1] / "out" / "rep_pipelines"
    if args.dry_run:
        out.mkdir(parents=True, exist_ok=True)
    results = {}
    master = build_master(by_rep, stats, generated)
    results[MASTER_NAME] = (out / MASTER_NAME).write_bytes(master) and "written" if args.dry_run else upload(MASTER_NAME, master, items)
    for rep, mine in by_rep.items():
        if rep == "Unattributed":
            continue  # no rep to share it with; the master keeps the sheet
        name = rep_filename(rep)
        data = build_rep_file(rep, mine, stats[rep], generated)
        results[name] = (out / name).write_bytes(data) and "written" if args.dry_run else upload(name, data, items)
    for name, res in results.items():
        log(f"  {res:8} {name}")
    log(f"done: {len(by_rep)} reps, {len(pipeline)} pipeline leads, "
        f"quoted £{sum(s['quoted'] for s in stats.values()):,.0f}, est £{sum(s['est'] for s in stats.values()):,.0f}")
    json.dump({rep: stats[rep] for rep in stats}, open(out / "last_run.json", "w") if args.dry_run else open(os.devnull, "w"), indent=1)


if __name__ == "__main__":
    main()
